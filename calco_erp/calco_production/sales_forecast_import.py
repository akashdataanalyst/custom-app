from __future__ import annotations
from calco_erp.calco_production.forecast_period import assign_period, effective_period, period_sql

import csv
import math
import re
from datetime import date
from pathlib import Path

import frappe
from frappe import _
from frappe.utils import cint, cstr, get_first_day, get_last_day, getdate
from openpyxl import load_workbook


SALES_FORECAST_DOCTYPE = "Sales Forecast"
FINISHED_GOODS_ITEM_GROUP = "Finished Goods"
FORECAST_UOM = "Kg"
MAX_IMPORT_ROWS = 5000
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
HEADER_ALIASES = {
    "item_code": {"item code", "fg code"},
    "demand_qty": {"demand qty", "forecast qty", "forecast quantity"},
}


def normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", cstr(value).strip()).casefold()


def parse_positive_quantity(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None

    text = cstr(value).strip()
    if not text:
        return None

    try:
        quantity = float(text)
    except (TypeError, ValueError):
        return None

    if not math.isfinite(quantity) or quantity <= 0:
        return None
    return quantity


def read_source_rows(path: str | Path) -> list[list[object]]:
    source_path = Path(path)
    extension = source_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        frappe.throw(_("Only .xlsx and .csv forecast files are supported."))

    if extension == ".xlsx":
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()

    with source_path.open("r", encoding="utf-8-sig", newline="") as source_file:
        return [list(row) for row in csv.reader(source_file)]


def resolve_columns(header: list[object]) -> dict[str, int]:
    resolved: dict[str, list[int]] = {fieldname: [] for fieldname in HEADER_ALIASES}
    unsupported_headers: list[str] = []

    for index, value in enumerate(header):
        normalized = normalize_header(value)
        if not normalized:
            continue

        matches = [
            fieldname
            for fieldname, aliases in HEADER_ALIASES.items()
            if normalized in aliases
        ]
        if matches:
            resolved[matches[0]].append(index)
        else:
            unsupported_headers.append(cstr(value).strip())

    errors = []
    for fieldname, indexes in resolved.items():
        label = "Item Code" if fieldname == "item_code" else "Demand Qty"
        if not indexes:
            errors.append(_("Missing required column: {0}").format(label))
        elif len(indexes) > 1:
            errors.append(_("Multiple columns map to {0}.").format(label))

    if unsupported_headers:
        errors.append(
            _("Unsupported columns: {0}").format(", ".join(unsupported_headers))
        )
    if errors:
        frappe.throw("<br>".join(errors), title=_("Invalid Forecast Template"))

    return {fieldname: indexes[0] for fieldname, indexes in resolved.items()}


def parse_forecast_rows(path: str | Path) -> list[dict[str, object]]:
    source_rows = read_source_rows(path)
    if not source_rows:
        frappe.throw(_("The forecast file is empty."))

    columns = resolve_columns(source_rows[0])
    parsed_rows: list[dict[str, object]] = []
    for row_number, source_row in enumerate(source_rows[1:], start=2):
        item_value = (
            source_row[columns["item_code"]]
            if columns["item_code"] < len(source_row)
            else None
        )
        quantity_value = (
            source_row[columns["demand_qty"]]
            if columns["demand_qty"] < len(source_row)
            else None
        )
        if not cstr(item_value).strip() and not cstr(quantity_value).strip():
            continue

        parsed_rows.append(
            {
                "row_number": row_number,
                "item_code": cstr(item_value).strip(),
                "demand_qty_raw": quantity_value,
            }
        )

    if not parsed_rows:
        frappe.throw(_("The forecast file contains no data rows."))
    if len(parsed_rows) > MAX_IMPORT_ROWS:
        frappe.throw(
            _("A maximum of {0} forecast rows may be imported at once.").format(
                MAX_IMPORT_ROWS
            )
        )
    return parsed_rows


def get_import_forecast_period(doc) -> tuple[date, date]:
    if not doc.company:
        frappe.throw(_("Company is required before importing a forecast."))
    if not doc.from_date:
        frappe.throw(_("From Date is required before importing a forecast."))
    if doc.frequency != "Monthly":
        frappe.throw(_("Forecast Excel import supports Monthly frequency only."))
    if cint(doc.demand_number) != 1:
        frappe.throw(_("Number of Months must be 1 for Forecast Excel import."))
    if not doc.parent_warehouse:
        frappe.throw(_("Parent Warehouse is required before importing a forecast."))

    forecast_date = getdate(doc.from_date)
    return get_first_day(forecast_date), get_last_day(forecast_date)


def get_finished_goods_warehouse(doc) -> str:
    warehouse = frappe.db.get_value("Company", doc.company, "default_fg_warehouse")
    if not warehouse:
        frappe.throw(
            _("Configure Default Finished Goods Warehouse for Company {0}.").format(
                doc.company
            )
        )

    warehouse_details = frappe.db.get_value(
        "Warehouse",
        warehouse,
        ["name", "company", "disabled", "is_group", "lft", "rgt"],
        as_dict=True,
    )
    parent_details = frappe.db.get_value(
        "Warehouse",
        doc.parent_warehouse,
        ["name", "company", "disabled", "is_group", "lft", "rgt"],
        as_dict=True,
    )
    if not warehouse_details:
        frappe.throw(_("Default Finished Goods Warehouse {0} does not exist.").format(warehouse))
    if warehouse_details.disabled:
        frappe.throw(_("Default Finished Goods Warehouse {0} is disabled.").format(warehouse))
    if warehouse_details.is_group:
        frappe.throw(_("Default Finished Goods Warehouse {0} cannot be a group warehouse.").format(warehouse))
    if warehouse_details.company != doc.company:
        frappe.throw(_("Default Finished Goods Warehouse {0} belongs to another company.").format(warehouse))
    if not parent_details or not parent_details.is_group or parent_details.disabled:
        frappe.throw(_("Select an enabled group warehouse as Parent Warehouse."))
    if parent_details.company != doc.company:
        frappe.throw(_("Parent Warehouse belongs to another company."))
    if not (
        cint(parent_details.lft) < cint(warehouse_details.lft)
        and cint(warehouse_details.rgt) < cint(parent_details.rgt)
    ):
        frappe.throw(
            _("Default Finished Goods Warehouse {0} is not below Parent Warehouse {1}.").format(
                warehouse,
                doc.parent_warehouse,
            )
        )
    return warehouse


def validate_and_normalize_rows(
    parsed_rows: list[dict[str, object]],
    delivery_date: date,
    warehouse: str,
) -> list[dict[str, object]]:
    normalized_rows: list[dict[str, object]] = []
    errors: list[str] = []
    seen_items: dict[str, int] = {}

    for source_row in parsed_rows:
        row_number = cint(source_row["row_number"])
        item_code = cstr(source_row.get("item_code")).strip()
        quantity = parse_positive_quantity(source_row.get("demand_qty_raw"))
        if not item_code:
            errors.append(_("Row {0}: Item Code is required.").format(row_number))
            continue
        if quantity is None:
            errors.append(
                _("Row {0}: Demand Qty must be numeric and greater than zero.").format(
                    row_number
                )
            )
            continue

        item = frappe.db.get_value(
            "Item",
            item_code,
            ["name", "item_name", "disabled", "item_group", "stock_uom"],
            as_dict=True,
        )
        if not item:
            errors.append(_("Row {0}: Item {1} does not exist.").format(row_number, item_code))
            continue
        if item.disabled:
            errors.append(_("Row {0}: Item {1} is disabled.").format(row_number, item.name))
            continue
        if item.item_group != FINISHED_GOODS_ITEM_GROUP:
            errors.append(_("Row {0}: Item {1} is not a Finished Good.").format(row_number, item.name))
            continue
        if item.stock_uom != FORECAST_UOM:
            errors.append(
                _("Row {0}: Item {1} Stock UOM must be Kg.").format(row_number, item.name)
            )
            continue

        identity = cstr(item.name).casefold()
        if identity in seen_items:
            errors.append(
                _("Row {0}: Item {1} duplicates row {2}.").format(
                    row_number,
                    item.name,
                    seen_items[identity],
                )
            )
            continue
        seen_items[identity] = row_number
        normalized_rows.append(
            {
                "row_number": row_number,
                "item_code": item.name,
                "item_name": item.item_name,
                "uom": FORECAST_UOM,
                "demand_qty": quantity,
                "delivery_date": delivery_date,
                "warehouse": warehouse,
            }
        )

    if errors:
        frappe.throw("<br>".join(errors), title=_("Forecast Validation Failed"))
    return normalized_rows


def get_source_file(doc, file_name: str):
    file_doc = frappe.get_doc("File", file_name)
    file_doc.check_permission("read")
    if not file_doc.is_private:
        frappe.throw(_("Forecast source files must be private."))
    if doc.is_new():
        if file_doc.attached_to_doctype or file_doc.attached_to_name:
            frappe.throw(_("The forecast file is already attached to another document."))
    elif (
        file_doc.attached_to_doctype != SALES_FORECAST_DOCTYPE
        or file_doc.attached_to_name != doc.name
    ):
        frappe.throw(_("The forecast file must be attached to this Sales Forecast."))
    if Path(file_doc.file_name).suffix.lower() not in SUPPORTED_EXTENSIONS:
        frappe.throw(_("Only .xlsx and .csv forecast files are supported."))
    return file_doc


def get_forecast_doc(forecast_name: str | None, forecast_doc: str | dict | None):
    if forecast_name:
        doc = frappe.get_doc(SALES_FORECAST_DOCTYPE, forecast_name)
        doc.check_permission("write")
        return doc

    values = frappe.parse_json(forecast_doc) if forecast_doc else None
    if not values or values.get("doctype") != SALES_FORECAST_DOCTYPE:
        frappe.throw(_("Sales Forecast header details are required before import."))
    if not frappe.has_permission(SALES_FORECAST_DOCTYPE, ptype="create"):
        frappe.throw(_("You do not have permission to create a Sales Forecast."))
    return frappe.get_doc(values)


def load_import_context(
    file_name: str,
    forecast_name: str | None = None,
    forecast_doc: str | dict | None = None,
):
    doc = get_forecast_doc(forecast_name, forecast_doc)
    if doc.docstatus != 0:
        frappe.throw(_("Forecast Excel import is allowed only for Draft Sales Forecasts."))

    forecast_month, delivery_date = get_import_forecast_period(doc)
    warehouse = get_finished_goods_warehouse(doc)
    file_doc = get_source_file(doc, file_name)
    parsed_rows = parse_forecast_rows(file_doc.get_full_path())
    normalized_rows = validate_and_normalize_rows(parsed_rows, delivery_date, warehouse)
    return doc, file_doc, forecast_month, delivery_date, warehouse, normalized_rows


@frappe.whitelist()
def preview_forecast_import(
    file_name: str,
    forecast_name: str | None = None,
    forecast_doc: str | dict | None = None,
) -> dict[str, object]:
    doc, file_doc, forecast_month, delivery_date, warehouse, rows = load_import_context(
        file_name,
        forecast_name,
        forecast_doc,
    )
    return {
        "file_name": file_doc.name,
        "forecast_month": forecast_month.strftime("%B %Y"),
        "delivery_date": delivery_date,
        "warehouse": warehouse,
        "row_count": len(rows),
        "has_existing_rows": bool(doc.items),
        "is_new": doc.is_new(),
        "rows": rows,
    }


def replace_forecast_rows(doc, rows: list[dict[str, object]]) -> None:
    doc.set("items", [])
    doc.set("selected_items", [])
    for row in rows:
        doc.append(
            "items",
            {
                "item_code": row["item_code"],
                "item_name": row["item_name"],
                "uom": FORECAST_UOM,
                "delivery_date": row["delivery_date"],
                "demand_qty": row["demand_qty"],
                "warehouse": row["warehouse"],
            },
        )
        doc.append(
            "selected_items",
            {
                "item_code": row["item_code"],
                "demand_qty": row["demand_qty"],
            },
        )


@frappe.whitelist()
def apply_forecast_import(
    file_name: str,
    forecast_name: str | None = None,
    forecast_doc: str | dict | None = None,
    expected_modified: str | None = None,
    replace_existing: int = 0,
) -> dict[str, object]:
    doc, file_doc, forecast_month, delivery_date, warehouse, rows = load_import_context(
        file_name,
        forecast_name,
        forecast_doc,
    )
    is_new = doc.is_new()
    if not is_new and expected_modified and cstr(doc.modified) != cstr(expected_modified):
        frappe.throw(_("Sales Forecast changed after preview. Reload and import again."))
    if doc.items and not cint(replace_existing):
        frappe.throw(_("This Sales Forecast already has rows. Confirm replacement to continue."))

    replace_forecast_rows(doc, rows)
    for row in doc.items:
        assign_period(doc, row, 0)
    if is_new:
        doc.insert()
        file_doc.attached_to_doctype = SALES_FORECAST_DOCTYPE
        file_doc.attached_to_name = doc.name
        file_doc.save(ignore_permissions=True)
    else:
        doc.save()
    return {
        "name": doc.name,
        "file_name": file_doc.name,
        "forecast_month": forecast_month.strftime("%B %Y"),
        "delivery_date": delivery_date,
        "warehouse": warehouse,
        "row_count": len(rows),
        "docstatus": doc.docstatus,
    }


def get_forecast_months(doc) -> list[tuple[date, date]]:
    months: set[tuple[int, int]] = set()
    for row in doc.get("items") or []:
        if effective_period(row):
            row_date = getdate(effective_period(row))
            months.add((row_date.year, row_date.month))

    if not months and doc.from_date:
        forecast_date = getdate(doc.from_date)
        months.add((forecast_date.year, forecast_date.month))

    return [
        (get_first_day(date(year, month, 1)), get_last_day(date(year, month, 1)))
        for year, month in sorted(months)
    ]


def validate_unique_submitted_forecast_month(doc, method=None) -> None:
    if not doc.company:
        return

    for month_start, month_end in get_forecast_months(doc):
        duplicate = frappe.db.sql(
            f"""
            select distinct sf.name
            from `tabSales Forecast` sf
            inner join `tabSales Forecast Item` sfi
                on sfi.parent = sf.name
                and sfi.parenttype = 'Sales Forecast'
                and sfi.parentfield = 'items'
            where sf.docstatus = 1
                and sf.company = %s
                and sf.name != %s
                and {period_sql()} between %s and %s
            limit 1
            """,
            (doc.company, doc.name or "", month_start, month_end),
            as_dict=True,
        )
        if duplicate:
            frappe.throw(
                _("Submitted Sales Forecast {0} already exists for {1}.").format(
                    duplicate[0].name,
                    month_start.strftime("%B %Y"),
                ),
                title=_("Duplicate Monthly Sales Forecast"),
            )
