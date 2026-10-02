from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path

import frappe
from frappe.utils import cint, cstr, getdate
from frappe.utils.file_manager import save_file
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from calco_erp.calco_production.mpds_master import (
    CURRENT_STATUS,
    IMPORTED_STATUS,
    MPDS_DOCTYPE,
    SUPERSEDED_STATUS,
    _applicability,
    get_authority_payload_hash,
    parse_specification,
)


EXPECTED_WORKBOOK_SHA256 = "AAAADAA2126FC8C3607992EB78B5A9A6F4F12CEFDF3B8C10891B46F4B71EEB66"
SOURCE_SHEET = "Line2"
PILOT_GRADE = "710C3031"
LINE_MAP = {
    "L-1": "Line 1",
    "L-2": "Line 2",
    "L-3": "Line 3",
    "L-5": "Line 5",
    "L-6": "Line 6",
}
CONTROL_COLUMNS = {"A", "B", "BO", "BP", "BQ", "CC", "DA", "DB"}
SYNTHETIC_MISSING_PROCESS_PARAMETERS = (
    ("air_knife_working", "Air Knife Working"),
    ("vent_condition", "Vent Condition"),
    ("all_magnet_cleaning", "All Magnet Cleaning"),
)


PROCESS_PARAMETER_MAP = {
    "screw configuration": ("screw_configuration", "Screw Configuration", "Mapped", ""),
    "lube oil temp": ("gear_box_oil_temperature", "Gear Box Oil Temperature", "Mapped", "Business-confirmed equivalent: Lube Oil Temp -> Gear Box Oil Temperature. Original source header preserved."),
    "tcu temp": ("tcu_temperature", "TCU Temperature", "Mapped", ""),
    "screen pack": ("screen_pack", "Screen Pack", "Mapped", ""),
    "die": ("die_hole", "Die Hole", "Mapped", "Business-confirmed equivalent: Die -> Die Hole. Original source header preserved."),
    "water bath tempr 1 (°c)": ("water_bath_temperature_1", "Water Bath Temperature 1", "Mapped", ""),
    "water bath tempr 2 (°c)": ("water_bath_temperature_2", "Water Bath Temperature 2", "Mapped", ""),
    "temp sc adopter": ("adapter_temperature", "Adapter Temperature", "Mapped", "Source spelling preserved as Temp SC ADOPTER."),
    "tem screen changer": ("screen_changer_temperature", "Screen Changer Temperature", "Mapped", "Source spelling preserved as Tem Screen Changer."),
    "temp die head": ("die_head_temperature", "Die Head Temperature", "Mapped", ""),
    "vacumm stuffer temp": ("vacuum_stuffer_temperature", "Vacuum Stuffer Temperature", "Mapped", "Source spelling preserved as Vacumm Stuffer Temp."),
    "screw rpm": ("extruder_rpm", "Extruder RPM", "Mapped", ""),
    "sf1 rpm": ("sf1_rpm", "SF1 RPM", "Mapped", "Business-confirmed direct F-PRD process parameter."),
    "sf2 rpm": ("sf2_rpm", "SF2 RPM", "Mapped", "Business-confirmed direct F-PRD process parameter."),
    "torque (%)": ("torque", "Torque", "Mapped", ""),
    "current in amp": ("current", "Current", "Mapped", ""),
    "melt temperature": ("melt_temperature", "Melt Temperature", "Mapped", ""),
    "melt pressure": ("melt_pressure", "Melt Pressure", "Mapped", ""),
    "specific energy ( kwh/kg)": ("specific_energy", "Specific Energy", "Mapped", ""),
    "vacumm": ("vacuum", "Vacuum", "Mapped", "Source spelling preserved as Vacumm."),
    "throughtput": ("throughput", "Throughput", "Mapped", "Source spelling preserved as Throughtput."),
    "palletizer speed": ("pelletizer_frequency", "Pelletizer Frequency", "Mapped", "Business-confirmed equivalent: Palletizer Speed -> Pelletizer Frequency. Original source header preserved."),
    "granuel size (mm)": ("granule_size", "Granule Size", "Mapped", "Source spelling preserved as Granuel Size."),
    "classifier inlet temp": ("classifier_inlet_temperature", "Classifier Inlet Temperature", "Mapped", ""),
    "classifier outlet temp": ("classifier_outlet_temperature", "Classifier Outlet Temperature", "Mapped", ""),
}


def _normalized_header(value) -> str:
    return re.sub(r"\s+", " ", cstr(value)).strip().casefold()


def _workbook_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _source_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _source_value_type(value) -> str:
    if value is None:
        return "Blank"
    if isinstance(value, datetime):
        return "Datetime"
    if isinstance(value, date):
        return "Date"
    if isinstance(value, bool):
        return "Boolean"
    if isinstance(value, (int, float)):
        return "Number"
    return "Text"


def _extract_unit(header: str) -> str:
    normalized = _normalized_header(header)
    explicit = {
        "torque (%)": "%",
        "current in amp": "Amp",
        "water bath tempr 1 (°c)": "°C",
        "water bath tempr 2 (°c)": "°C",
        "fg packing material tempr (°c)": "°C",
        "specific energy ( kwh/kg)": "KWH/KG",
        "granuel size (mm)": "mm",
        "screw speed - f4 (kg/hr)": "kg/hr",
        "mixing time (min)": "min",
        "packing size (kg)": "kg",
    }
    return explicit.get(normalized, "")


def _domain_for(column: str, header: str) -> str:
    index = _column_number(column)
    if 3 <= index <= 40 or 85 <= index <= 91:
        return "Feeder Configuration"
    if 41 <= index <= 63 or 72 <= index <= 79:
        return "Extruder / Process Parameters"
    if 64 <= index <= 66:
        return "Water Bath / Cooling"
    if 80 <= index <= 84:
        return "Pelletizer / Classifier"
    if 70 <= index <= 71 or 86 <= index <= 95:
        return "Premix / Blending Reference"
    if 96 <= index <= 104:
        return "Packing / Downstream"
    return "Other Operational Instructions"


def _column_number(column: str) -> int:
    number = 0
    for character in column:
        number = number * 26 + ord(character.upper()) - 64
    return number


def _position_for(header: str) -> str:
    feeder = re.search(r"-\s*F([1-6])(?:\s|$)", cstr(header), re.IGNORECASE)
    if feeder:
        return f"F{feeder.group(1)}"
    zone = re.search(r"\bZ(\d+)\b", cstr(header), re.IGNORECASE)
    if zone:
        return f"Z{zone.group(1)}"
    return ""


def _process_mapping(header: str):
    normalized = _normalized_header(header)
    if normalized.startswith("temp z"):
        zone = re.search(r"z(\d+)", normalized)
        if zone:
            return (f"temperature_zone_{zone.group(1)}", f"Temperature Zone {zone.group(1)}", "Mapped", "")
    return PROCESS_PARAMETER_MAP.get(normalized)


def _specification_row(column: str, row_number: int, header, value, sequence: int) -> dict:
    header_text = cstr(header)
    value_text = _source_text(value)
    process_mapping = _process_mapping(header_text)
    display_label = process_mapping[1] if process_mapping else header_text
    mapping_status = process_mapping[2] if process_mapping else "Unmapped"
    mapping_note = process_mapping[3] if process_mapping else ""
    parsed = parse_specification(value_text)
    if _source_value_type(value) in {"Date", "Datetime"}:
        parsed["parser_status"] = "Warning"
        mapping_note = (mapping_note + " Source cell is date-formatted; value preserved without correction.").strip()
    return {
        "sequence": sequence,
        "domain": _domain_for(column, header_text),
        "parameter_key": f"{frappe.scrub(header_text) or 'unnamed'}__{column.casefold()}",
        "exact_source_header": header_text,
        "display_label": display_label,
        "feeder_or_position": _position_for(header_text),
        "specification_text": value_text,
        "unit": _extract_unit(header_text),
        "applicability": _applicability(value_text),
        "parameter_group": "F-PRD-01/01" if process_mapping else "MPDS",
        "process_parameter_key": process_mapping[0] if process_mapping else "",
        "mapping_status": mapping_status,
        "mapping_note": mapping_note,
        "source_column": column,
        "source_cell": f"{SOURCE_SHEET}!{column}{row_number}",
        "source_value_type": _source_value_type(value),
        "source_raw_value": value_text,
        **parsed,
    }


def _missing_process_row(parameter_key: str, label: str, row_number: int, sequence: int) -> dict:
    return {
        "sequence": sequence,
        "domain": "Extruder / Process Parameters",
        "parameter_key": f"missing_line2__{parameter_key}",
        "exact_source_header": "",
        "display_label": label,
        "feeder_or_position": "",
        "specification_text": "",
        "unit": "",
        "applicability": "Blank",
        "parameter_group": "F-PRD-01/01",
        "process_parameter_key": parameter_key,
        "mapping_status": "Missing Source",
        "mapping_note": "No authoritative Line2 source column exists. MPDS_2 static/fallback value was not imported.",
        "parser_status": "Blank",
        "source_column": "",
        "source_cell": "",
        "source_value_type": "Unavailable",
        "source_raw_value": "",
    }


def _identity_key(grade: str, line_identifier: str, source_row: int) -> str:
    raw = f"{SOURCE_SHEET}|{source_row}|{grade}|{line_identifier}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _source_payload_hash(payload: dict) -> str:
    source = {
        "fg_item": payload["fg_item"],
        "source_line_identifier": payload["source_line_identifier"],
        "source_mpds_no": payload["source_mpds_no"],
        "source_revision_text": payload["source_revision_text"],
        "source_revision_date": str(payload.get("source_revision_date") or ""),
        "source_row": payload["source_row"],
        "specifications": [
            {
                "source_cell": row["source_cell"],
                "exact_source_header": row["exact_source_header"],
                "source_value_type": row["source_value_type"],
                "source_raw_value": row["source_raw_value"],
            }
            for row in payload["specifications"]
        ],
    }
    return hashlib.sha256(json.dumps(source, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()


def _build_payload(sheet, row_number: int, workbook_hash: str) -> dict:
    values = {get_column_letter(column): sheet.cell(row_number, column).value for column in range(1, sheet.max_column + 1)}
    grade = _source_text(values["A"])
    line_identifier = _source_text(values["B"])
    production_line = LINE_MAP.get(line_identifier)
    if not production_line:
        raise ValueError(f"Unsupported authoritative source line {line_identifier!r} at {SOURCE_SHEET}!B{row_number}")
    specifications = []
    for column_number in range(3, sheet.max_column + 1):
        column = get_column_letter(column_number)
        if column in CONTROL_COLUMNS:
            continue
        header = sheet.cell(1, column_number).value
        if header is None:
            continue
        specifications.append(_specification_row(column, row_number, header, values[column], len(specifications) + 1))
    for parameter_key, label in SYNTHETIC_MISSING_PROCESS_PARAMETERS:
        specifications.append(_missing_process_row(parameter_key, label, row_number, len(specifications) + 1))
    revision_date = values["BQ"]
    if isinstance(revision_date, datetime):
        revision_date = revision_date.date()
    elif revision_date:
        revision_date = getdate(revision_date)
    source_revision = _source_text(values["BP"])
    source_mpds = _source_text(values["BO"])
    payload = {
        "fg_item": grade,
        "production_line": production_line,
        "source_line_identifier": line_identifier,
        "mpds_no": source_mpds,
        "revision": source_revision,
        "source_mpds_no": source_mpds,
        "source_revision_text": source_revision,
        "source_revision_date": revision_date,
        "effective_date": None,
        "material_status_snapshot": _source_text(values["DA"]),
        "source_workbook_hash": workbook_hash,
        "source_sheet": SOURCE_SHEET,
        "source_row": row_number,
        "import_identity_key": _identity_key(grade, line_identifier, row_number),
        "specifications": specifications,
    }
    payload["source_payload_hash"] = _source_payload_hash(payload)
    return payload


def inspect_pilot_workbook(source_path: str) -> dict:
    path = Path(source_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"MPDS workbook not found: {path}")
    workbook_hash = _workbook_hash(path)
    if workbook_hash != EXPECTED_WORKBOOK_SHA256:
        raise ValueError(f"MPDS workbook SHA-256 mismatch: expected {EXPECTED_WORKBOOK_SHA256}, received {workbook_hash}")
    workbook = load_workbook(path, data_only=True, read_only=False)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Authoritative sheet {SOURCE_SHEET} not found")
    sheet = workbook[SOURCE_SHEET]
    source_rows = [row for row in range(2, sheet.max_row + 1) if _source_text(sheet.cell(row, 1).value) == PILOT_GRADE]
    payloads = [_build_payload(sheet, row, workbook_hash) for row in source_rows]
    by_line = {}
    duplicates = []
    for payload in payloads:
        line = payload["source_line_identifier"]
        if line in by_line:
            duplicates.append(line)
            payload["duplicate_candidate"] = 1
            by_line[line]["duplicate_candidate"] = 1
        else:
            payload["duplicate_candidate"] = 0
            by_line[line] = payload
    return {
        "source_path": str(path),
        "workbook_sha256": workbook_hash,
        "sheet": SOURCE_SHEET,
        "grade": PILOT_GRADE,
        "source_rows": source_rows,
        "lines": [payload["source_line_identifier"] for payload in payloads],
        "duplicates": sorted(set(duplicates)),
        "payloads": payloads,
    }


def _preserve_private_source(path: Path, workbook_hash: str):
    file_name = f"MPDS_{workbook_hash[:12]}.xlsx"
    existing = frappe.db.get_value(
        "File", {"file_name": file_name, "is_private": 1}, "name"
    )
    if existing:
        return existing
    file_doc = save_file(file_name, path.read_bytes(), None, None, is_private=1)
    return file_doc.name


def _summarize_payload(payload: dict) -> dict:
    values = [row["specification_text"] for row in payload["specifications"]]
    return {
        "source_row": payload["source_row"],
        "source_line": payload["source_line_identifier"],
        "production_line": payload["production_line"],
        "mpds_no": payload["mpds_no"],
        "revision": payload["revision"],
        "source_revision_date": str(payload.get("source_revision_date") or ""),
        "specification_rows": len(payload["specifications"]),
        "blank_specifications": sum(value == "" for value in values),
        "explicit_dash": sum(value.strip() == "-" for value in values),
        "not_applicable": sum(value.strip().upper() in {"NA", "N/A"} for value in values),
        "source_payload_hash": payload["source_payload_hash"],
    }


def import_710c3031_pilot(source_path: str, dry_run: bool = True) -> dict:
    audit = inspect_pilot_workbook(source_path)
    result = {
        "dry_run": bool(cint(dry_run)),
        "workbook_sha256": audit["workbook_sha256"],
        "sheet": SOURCE_SHEET,
        "grade": PILOT_GRADE,
        "created": [],
        "unchanged": [],
        "skipped": [],
        "exceptions": [],
        "source_records": [_summarize_payload(payload) for payload in audit["payloads"]],
    }
    if audit["duplicates"]:
        result["exceptions"].append({"type": "Duplicate Grade + Line", "lines": audit["duplicates"]})
    if bool(cint(dry_run)):
        return result
    source_file = _preserve_private_source(Path(audit["source_path"]), audit["workbook_sha256"])
    for payload in audit["payloads"]:
        existing_name = frappe.db.get_value(MPDS_DOCTYPE, {"import_identity_key": payload["import_identity_key"]}, "name")
        if existing_name:
            existing = frappe.get_doc(MPDS_DOCTYPE, existing_name)
            if existing.status in {CURRENT_STATUS, SUPERSEDED_STATUS}:
                result["skipped"].append({"name": existing.name, "reason": f"Controlled status {existing.status}"})
                continue
            if existing.get("reviewed_on"):
                result["skipped"].append({"name": existing.name, "reason": "ERP record has completed technical review"})
                continue
            current_hash = get_authority_payload_hash(existing)
            if existing.imported_payload_hash and current_hash != existing.imported_payload_hash:
                result["skipped"].append({"name": existing.name, "reason": "ERP record was reviewed or edited after import"})
                continue
            if existing.source_payload_hash == payload["source_payload_hash"]:
                result["unchanged"].append(existing.name)
                continue
            result["skipped"].append({"name": existing.name, "reason": "Source changed; explicit reconciliation required"})
            continue
        payload["source_file"] = source_file
        doc = frappe.get_doc({"doctype": MPDS_DOCTYPE, **payload})
        frappe.flags.mpds_import = True
        try:
            doc.insert()
        finally:
            frappe.flags.mpds_import = False
        stored = frappe.get_doc(MPDS_DOCTYPE, doc.name)
        payload_hash = get_authority_payload_hash(stored)
        frappe.db.set_value(MPDS_DOCTYPE, doc.name, "imported_payload_hash", payload_hash, update_modified=False)
        result["created"].append(doc.name)
    return result
