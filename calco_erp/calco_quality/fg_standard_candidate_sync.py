"""Opt-in FG Standard-only successor import; no Control Plan/MPDS routing."""
from __future__ import annotations
import hashlib,json
from collections import Counter,defaultdict
from datetime import date,datetime
from pathlib import Path
from openpyxl import load_workbook
import frappe
from frappe.utils import now_datetime
from calco_erp.calco_quality import quality_master_versions as q
VERSION='fg-standard-header-authority-v2'
META={'Product Name','Code','Rev. No.','Issue Date','Material Status_1'}
NA={'n/a','na','not applicable'}

def text(v):return '' if v is None else str(v).strip()

def payload(headers,values):
 from calco_erp.calco_quality.importers.fg_standard_min_max_import import parse_standard_header
 from calco_erp.calco_quality.fg_quality_setup import get_true_fg_testing_method_name_set
 props={};review=[];states=[];cells=q.source_cells({'headers':headers,'values':values})
 for i,(header,value) in enumerate(zip(headers,values)):
  if header in META:continue
  h=' '.join(text(header).split());v=text(value)
  state='Blank' if value is None or value=='' else 'Explicit Not Applicable' if v.casefold() in NA else 'Unspecified placeholder' if v=='-' else 'Invalid' if isinstance(value,(date,datetime)) or v.startswith('=') else 'Explicit specification'
  states.append({'column':i+1,'state':state})
  if state=='Invalid':review.append({'column':i+1,'header':header,'reason':'Date/formula cannot be used as specification authority'});continue
  if 'thickness' in h.casefold():continue
  mapped=parse_standard_header('Glow Wire Test @ 3.2 mm' if h.casefold()=='glow wire test @ 3.2mm (°c)' else h,get_true_fg_testing_method_name_set())
  if h=='Surface Appearance':mapped={'parameter':'Surface Appearance','fieldname':'target_value','unit':''}
  if not mapped:
   if state=='Explicit specification':review.append({'column':i+1,'header':header,'reason':'Preserved exact source property; no approved Calco parameter alias'})
   continue
  name=mapped['parameter'];prop=props.setdefault(name,{'parameter':name,'minimum_value':None,'maximum_value':None,'target_values':[],'unit':mapped.get('unit',''),'source_columns':[],'conditions':[],'not_applicable':False,'conflicts':[],'source_states':[]})
  prop['source_columns'].append(i+1);prop['source_states'].append(state)
  if state=='Explicit Not Applicable':prop['not_applicable']=True;continue
  if state!='Explicit specification':continue
  field=mapped['fieldname']
  if field in ('minimum_value','maximum_value') and q.numeric(value):
   if prop[field] is not None and prop[field]!=value:prop['conflicts'].append('Competing source '+field)
   else:prop[field]=value
  elif v not in prop['target_values']:prop['target_values'].append(v)
 for h,param in [('Flammability\n(thickness (mm))','UL94 Burning Testing'),('Glow Wire Test\nthickness (mm)','Glow Wire Test')]:
  if h in headers and (value:=values[headers.index(h)]) is not None and param in props:props[param]['conditions'].append('Source thickness (mm): '+text(value))
 for name,prop in props.items():
  lo,hi=prop['minimum_value'],prop['maximum_value']
  if lo is not None and hi is not None and lo>hi:prop['conflicts'].append('Minimum exceeds maximum')
  if prop['not_applicable'] and (lo is not None or hi is not None or prop['target_values']):prop['conflicts'].append('Explicit N/A conflicts with a specification in the same parameter')
  if prop['conflicts']:review.append({'parameter':name,'reason':'; '.join(prop['conflicts'])})
 return {'kind':'FG Standard','cells':cells,'properties':props,'property_review':review,'cell_states':states,'parser_version':VERSION,'ipqc_timing':'Not supplied by FG Standard'}

def inspect(path,expected_sha256):
 raw=Path(path).read_bytes()
 if hashlib.sha256(raw).hexdigest()!=expected_sha256:raise ValueError('Candidate source fingerprint mismatch')
 w=load_workbook(path,read_only=True,data_only=False)
 try:
  if w.sheetnames!=['Sheet1']:raise ValueError('Authoritative sheet requires review')
  rows=list(w['Sheet1'].values);headers=list(rows[0]);idx={text(h):i for i,h in enumerate(headers)}
  if not META.issubset(idx):raise ValueError('Required source headers missing')
  groups=defaultdict(list)
  for number,row in enumerate(rows[1:],2):
   if text(row[idx['Code']]):groups[text(row[idx['Code']])].append((number,list(row)))
  items={r.name:r for r in frappe.get_all('Item',fields=['name','disabled'],limit_page_length=0)}
  result=[]
  for fg,group in groups.items():
   old=q.current('FG Standard',fg)
   for number,values in group:
    data=payload(headers,values);status=text(values[idx['Material Status_1']]);key='QMV-FG2-'+q.source.fingerprint([VERSION,expected_sha256,number,data])[:32]
    issues=[]
    if fg not in items:issues.append('Missing FG Item')
    elif items[fg].disabled:issues.append('Disabled FG Item')
    if len(group)>1:issues.append('Source duplicate: competing revision/status')
    issues += [r['reason'] for r in data['property_review'] if r['reason']!='Preserved exact source property; no approved Calco parameter alias']
    classification='Held' if issues else 'Unchanged' if frappe.db.exists(q.DOCTYPE,key) else 'New Revision' if old else 'New FG Standard'
    result.append({'fg':fg,'row':number,'status':status,'revision':text(values[idx['Rev. No.']]),'issue_date':text(values[idx['Issue Date']]),'key':key,'old':old.name if old else None,'payload':data,'classification':classification,'issues':issues})
  return result
 finally:w.close()

def sync(path,expected_sha256,dry_run=True,expected_site=None,approval_reference=None,legacy_adoptions=None):
 if not isinstance(dry_run,bool):raise ValueError('dry_run must be boolean')
 if not dry_run:
  if expected_site!=frappe.local.site or not approval_reference:raise ValueError('Explicit site and approval reference required')
  if frappe.session.user!='Administrator' and 'Quality Manager' not in frappe.get_roles():frappe.throw('Quality Manager authority required',frappe.PermissionError)
 records=inspect(path,expected_sha256)
 for r in records:
  if dry_run or r['classification'] not in ('New Revision','New FG Standard'):continue
  frappe.db.sql('select name from `tabItem` where name=%s for update',(r['fg'],))
  if frappe.db.exists(q.DOCTYPE,r['key']):r['classification']='Unchanged';continue
  old=q.current('FG Standard',r['fg'])
  if (old.name if old else None)!=r['old']:frappe.throw('FG Standard changed concurrently; retry preflight')
  flag=getattr(frappe.flags,'controlled_quality_master_sync',False);frappe.flags.controlled_quality_master_sync=True
  try:
   if old:old.is_active=0;old.save()
   frappe.get_doc({'doctype':q.DOCTYPE,'version_key':r['key'],'master_kind':'FG Standard','fg_item':r['fg'],'production_line':'','is_active':1,'business_revision':r['revision'],'issue_date_text':r['issue_date'],'material_status':r['status'],'source_file':'FG Standard Updated.xlsx','source_sha256':expected_sha256,'source_sheet':'Sheet1','source_row':r['row'],'source_version':VERSION,'source_fingerprint':q.source.fingerprint([expected_sha256,r['row']]),'payload_hash':q.source.fingerprint(r['payload']),'payload':q.encoded(r['payload']),'approval_reference':approval_reference,'imported_by':frappe.session.user,'imported_on':now_datetime(),'supersedes':r['old'],'review_notes':q.encoded(dict(legacy_adoptions[r['key']],created_on=str(now_datetime()))) if legacy_adoptions and r['key'] in legacy_adoptions else 'FG-only header-driven source revision; no other master or transaction changed.'}).insert()
  finally:frappe.flags.controlled_quality_master_sync=flag
 return {'counts':dict(Counter(r['classification'] for r in records)),'records':records}
